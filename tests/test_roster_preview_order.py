"""Regression fixtures mirror table structure only; all student data is synthetic."""
import copy
import csv
from io import BytesIO, StringIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import SESSIONS, app, dataset_from_rows, recognize_header

client = TestClient(app)


@pytest.mark.parametrize('header,key', [
    ('Student name', 'student_name'), ('student name', 'student_name'),
    ('Student names', 'student_name'), ('Ism/familiya', 'student_name'),
    ('Ism va familiya', 'student_name'), ('ФИО', 'student_name'),
    ('F.I.O.', 'student_name'), ('№', 'row_number'), ('T/N', 'row_number'),
    ('T/R', 'row_number'), ('Parent name', 'parent_name'), ('ФИО родителя', 'parent_name'),
    ('Current notes', 'custom_currentnotes'), ('Birth date', 'birth_date'),
])
def test_headers_do_not_match_letters_inside_other_words(header, key):
    assert recognize_header(header).key == key


def synthetic_workbook(count=135):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['№', 'Student name', 'Birth date', 'JSHSHIR', 'Class', "sinfiga ko‘chirildi"])
    for i in range(count):
        if i % 25 == 0:
            sheet.append(['Telefon', 'Elektron pochta', None, None, None, '4-A'])
        sheet.append([i % 25 + 1, f'Example Learner {i + 1:03}', '01/01/2015', f'000000000{i:05}', '4-B', '4-A'])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture
def roster():
    uploaded = client.post('/api/upload', files=[('files', ('synthetic.xlsx', synthetic_workbook(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))])
    assert uploaded.status_code == 200, uploaded.text
    data = uploaded.json()
    assert data['row_count'] == 135
    assert {'row_number', 'student_name', 'birth_date', 'pinfl', 'class', 'promoted_class'} <= {col['key'] for col in data['columns']}
    request = {
        'session_id': data['session_id'], 'dataset_id': data['active_dataset_id'],
        'columns': [{'key': key, 'name': name} for key, name in [('row_number', '№'), ('student_name', 'Name'), ('pinfl', 'ID'), ('class', 'Class'), ('promoted_class', 'Destination')]],
        'derived_columns': [{'name': 'Sequence', 'kind': 'sequence', 'start': 5}, {'name': 'Next', 'kind': 'next_class', 'source_key': 'class'}, {'name': 'Number copy', 'kind': 'copy', 'source_key': 'row_number'}],
        'manual_rows': [{'student_name': 'Added Learner', 'class': '7-A', 'pinfl': '0000042'}],
        'manual_row_ids': ['manual:added'],
    }
    yield request
    client.delete('/api/session/' + data['session_id'])


def test_value_evidence_recovers_duplicate_number_heading_without_losing_values():
    dataset = dataset_from_rows([
        ['No.', 'No.', 'Class'],
        [1, 'Example Learner Alpha', '4-A'],
        [2, 'Example Learner Beta', '4-A'],
        [3, 'Example Learner Gamma', '4-A'],
    ], 'synthetic.csv', 'Table')
    assert dataset is not None
    assert {c.key for c in dataset['columns']} == {'row_number', 'student_name', 'class'}
    assert dataset['records'][0]['student_name'] == 'Example Learner Alpha'
    assert dataset['records'][0]['row_number'] == 1


def test_unknown_numbers_are_not_assumed_to_be_row_numbers():
    dataset = dataset_from_rows([
        ['Reference', 'Student name', 'Age'],
        [452, 'Example Learner Alpha', 10],
        [197, 'Example Learner Beta', 10],
        [823, 'Example Learner Gamma', 11],
    ], 'synthetic.csv', 'Table')
    keys = {c.key for c in dataset['columns']}
    assert 'row_number' not in keys
    assert {'custom_reference', 'student_name', 'age'} <= keys


def test_pagination_reaches_every_row_and_focuses_additions(roster):
    pages = [client.post('/api/preview', json={**roster, 'offset': offset, 'limit': 50}).json() for offset in (0, 50, 100)]
    assert [page['preview_count'] for page in pages] == [50, 50, 36]
    assert len({row_id for page in pages for row_id in page['row_ids']}) == 136
    assert pages[-1]['rows'][-1]['Name'] == 'Added Learner'
    focused = client.post('/api/preview', json={**roster, 'limit': 50, 'focus_row_id': 'manual:added'}).json()
    assert focused['offset'] == 100 and focused['focus_found']
    clamped = client.post('/api/preview', json={**roster, 'offset': 999999, 'limit': 50}).json()
    assert clamped['offset'] == 100


def test_reordering_renumbers_and_recalculates_without_mutating_source(roster):
    original = copy.deepcopy(SESSIONS[roster['session_id']]['records'])
    initial = client.post('/api/preview', json=roster).json()
    ids = initial['row_order']
    order = [ids[2], ids[-1], ids[0], ids[1], *ids[3:-1]]
    request = {**roster, 'row_order': order, 'limit': 200}
    result = client.post('/api/preview', json=request).json()
    assert result['row_order'] == order
    rows = result['rows']
    assert [row['Name'] for row in rows[:4]] == ['Example Learner 003', 'Added Learner', 'Example Learner 001', 'Example Learner 002']
    assert [row['№'] for row in rows] == list(range(1, 137))
    assert [row['Number copy'] for row in rows] == list(range(1, 137))
    assert [row['Sequence'] for row in rows] == list(range(5, 141))
    assert rows[1]['Next'] == '8-A' and rows[1]['ID'] == '0000042'
    assert rows[0]['Destination'] == '4-A'  # Supplied destination stays with its student.
    assert SESSIONS[roster['session_id']]['records'] == original


@pytest.mark.parametrize('format', ['csv', 'xlsx', 'docx'])
def test_exports_match_reordered_preview_including_additions(roster, format):
    request = {**roster, 'row_order': ['manual:added', 'source:2', 'source:0', 'source:1'], 'limit': 200}
    preview = client.post('/api/preview', json=request).json()
    exported = client.post('/api/export', json={**request, 'format': format})
    assert exported.status_code == 200, exported.text
    expected = [[str(row[header]) for header in preview['headers']] for row in preview['rows']]
    if format == 'csv':
        actual = list(csv.reader(StringIO(exported.content.decode('utf-8-sig'))))[1:]
    elif format == 'xlsx':
        actual = [[str(value) if value is not None else '' for value in row] for row in list(load_workbook(BytesIO(exported.content)).active.values)[1:]]
    else:
        actual = [[cell.text for cell in row.cells] for row in Document(BytesIO(exported.content)).tables[0].rows[1:]]
    assert actual == expected


def test_filter_deduplication_and_sort_leave_no_numbering_gaps(roster):
    roster['manual_rows'] = [{'student_name': 'Example Learner 001', 'class': '4-B', 'pinfl': '00000000000000', 'promoted_class': '4-A'}, {'class': '4-A'}, {'student_name': 'Added Learner', 'class': '7-A'}]
    roster['manual_row_ids'] = ['manual:duplicate', 'manual:blank', 'manual:kept']
    roster['options'] = {'skip_blank_key': 'student_name', 'remove_duplicates': True, 'sort_key': 'student_name'}
    response = client.post('/api/preview', json={**roster, 'limit': 200}).json()
    assert response['row_count'] == 136
    assert [r['№'] for r in response['rows']] == list(range(1, 137))
    assert [r['Sequence'] for r in response['rows']] == list(range(5, 141))
    focused = client.post('/api/preview', json={**roster, 'focus_row_id': 'manual:blank'}).json()
    assert not focused['focus_found']


@pytest.mark.parametrize('change', [
    {'row_order': ['source:999999']}, {'row_order': ['source:0', 'source:0']},
    {'manual_row_ids': ['manual:a', 'manual:b']}, {'row_order': ['manual:missing']},
])
def test_invalid_identity_cannot_duplicate_or_inject_students(roster, change):
    response = client.post('/api/preview', json={**roster, **change})
    assert response.status_code == 400


def test_output_label_collision_does_not_silently_overwrite_column(roster):
    roster['columns'][1]['name'] = '№'
    assert client.post('/api/preview', json=roster).status_code == 400


def test_clients_can_explicitly_preserve_source_numbering(roster):
    response = client.post('/api/preview', json={**roster, 'options': {'renumber_rows': False}, 'row_order': ['source:2', 'source:0']}).json()
    assert [row['№'] for row in response['rows'][:2]] == [3, 1]
