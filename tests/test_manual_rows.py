"""Synthetic students only; exercise the real upload/preview/export API."""
import copy
import csv
from io import BytesIO, StringIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import SESSIONS, app

client = TestClient(app)


@pytest.fixture
def roster():
    response = client.post('/api/upload', files=[
        ('files', ('roster.csv', b'F.I.O.,Class,Phone\nDemo Learner,8-A,00123\n', 'text/csv')),
    ])
    assert response.status_code == 200, response.text
    data = response.json()
    assert data['active_dataset_id'] is not None
    assert {'student_name', 'class', 'phone'} <= {column['key'] for column in data['columns']}
    request = {
        'session_id': data['session_id'],
        'dataset_id': data['active_dataset_id'],
        'columns': [
            {'key': 'student_name', 'name': 'Name'},
            {'key': 'class', 'name': 'Class'},
            {'key': 'phone', 'name': 'Phone'},
        ],
        'fixed_columns': [{'name': 'Year', 'value': '2026–2027'}],
        'derived_columns': [
            {'name': 'Next', 'kind': 'next_class', 'source_key': 'class'},
            {'name': 'No.', 'kind': 'sequence'},
        ],
        'manual_rows': [{'student_name': '  Sample   Learner  ', 'class': '7-B', 'phone': '00042'}],
    }
    yield request
    client.delete('/api/session/' + data['session_id'])


def test_preview_appends_and_recalculates_without_mutating_source(roster):
    original = copy.deepcopy(SESSIONS[roster['session_id']]['records'])
    for _ in range(2):  # Preview/export requests must not append repeatedly to the session.
        response = client.post('/api/preview', json={**roster, 'limit': 1})
        assert response.status_code == 200, response.text
        assert response.json()['row_count'] == 2
        assert response.json()['preview_count'] == 1
    rows = client.post('/api/preview', json=roster).json()['rows']
    assert rows[-1] == {'Name': 'Sample Learner', 'Class': '7-B', 'Phone': '00042',
                        'Year': '2026–2027', 'Next': '8-B', 'No.': 2}
    assert SESSIONS[roster['session_id']]['records'] == original
    # Editing/removing additions replaces the request's draft, not stored source rows.
    roster['manual_rows'][0]['student_name'] = 'Edited Learner'
    assert client.post('/api/preview', json=roster).json()['rows'][-1]['Name'] == 'Edited Learner'
    roster['manual_rows'] = []
    assert client.post('/api/preview', json=roster).json()['row_count'] == 1


@pytest.mark.parametrize('format', ['xlsx', 'csv', 'docx'])
def test_all_downloads_include_manual_student(roster, format):
    response = client.post('/api/export', json={**roster, 'format': format})
    assert response.status_code == 200, response.text
    if format == 'xlsx':
        sheet = load_workbook(BytesIO(response.content), data_only=True).active
        rows = list(sheet.values)
    elif format == 'csv':
        rows = list(csv.reader(StringIO(response.content.decode('utf-8-sig'))))
    else:
        rows = [[cell.text for cell in row.cells]
                for row in Document(BytesIO(response.content)).tables[0].rows]
    assert len(rows) == 3
    assert list(rows[-1])[:5] == ['Sample Learner', '7-B', '00042', '2026–2027', '8-B']
    assert str(rows[-1][-1]) == '2'


def test_blank_additions_and_existing_cleaning_rules(roster):
    roster['manual_rows'] += [{}, {'student_name': ' \t\n '}, {'class': '5-A'},
                              {'student_name': 'Demo Learner', 'class': '8-A', 'phone': '00123'}]
    roster['derived_columns'] = []
    roster['options'] = {'skip_blank_key': 'student_name', 'remove_duplicates': True, 'sort_key': 'class'}
    response = client.post('/api/preview', json=roster)
    assert response.status_code == 200, response.text
    assert [row['Name'] for row in response.json()['rows']] == ['Sample Learner', 'Demo Learner']


def test_explicit_dataset_isolation_even_after_selection_changes(roster):
    session = SESSIONS[roster['session_id']]
    other = copy.deepcopy(session['dataset_groups'][roster['dataset_id']])
    other.update(id='second', records=[{'student_name': 'Other Student', 'class': '2-A'}], row_count=1)
    session['dataset_groups']['second'] = other
    result = client.post('/api/select-dataset', json={'session_id': roster['session_id'], 'dataset_id': 'second'})
    assert result.status_code == 200
    first = client.post('/api/preview', json=roster).json()
    assert [row['Name'] for row in first['rows']] == ['Demo Learner', 'Sample Learner']
    second = client.post('/api/preview', json={**roster, 'dataset_id': 'second', 'manual_rows': []}).json()
    assert [row['Name'] for row in second['rows']] == ['Other Student']
    assert session['active_dataset_id'] == 'second'


@pytest.mark.parametrize('changes,status', [
    ({'dataset_id': 'missing'}, 404),
    ({'dataset_id': None}, 400),
    ({'manual_rows': [{'not_a_column': 'value'}]}, 400),
    ({'manual_rows': [{}] * 101}, 422),
    ({'manual_rows': [{'student_name': 'a' * 2001}]}, 422),
    ({'manual_rows': [{'student_name': {'nested': 'bad'}}]}, 422),
    ({'manual_rows': [{str(i): '' for i in range(201)}]}, 422),
    ({'manual_rows': [{'student_name': 'a' * 2000, 'class': 'b' * 2000}] * 51}, 422),
])
def test_invalid_manual_rows_fail_with_clear_http_errors(roster, changes, status):
    for endpoint in ('preview', 'export'):
        response = client.post('/api/' + endpoint, json={**roster, **changes, 'format': 'xlsx'})
        assert response.status_code == status, response.text


def test_legacy_client_and_expired_session(roster):
    del roster['dataset_id']
    del roster['manual_rows']
    assert client.post('/api/preview', json=roster).json()['row_count'] == 1
    client.delete('/api/session/' + roster['session_id'])
    assert client.post('/api/preview', json=roster).status_code == 404
